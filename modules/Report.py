import os

os.environ["GIO_USE_VFS"] = "local"
os.environ["GDK_BACKEND"] = "none"
os.environ["NO_AT_BRIDGE"] = "1"
os.environ["WEASYPRINT_GUI"] = "false"

import pandas as pd
import base64
from datetime import datetime
import pdfkit
from weasyprint import HTML, CSS

import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
import platform


# ==========================================================
# 🔹 Utility Functions
# ==========================================================

def check(set_wg, act_wg, tol):
    if set_wg - tol <= act_wg <= set_wg + tol:
        return "Good"

    return "Up" if act_wg > set_wg + tol else "Down"


def difference(set_wg, act_wg):
    return round(act_wg - set_wg, 2)


_logo_cache = {}
LOGO_MAX_HEIGHT = 300   # px - plenty for the report header, at print resolution


def encode_logo(logo_path):
    """Logo as Base64 PNG for embedding, scaled down to print size.

    The uploaded logo can be huge (12500 x 3125 px, 2.6 MB); decoding that for
    every report page made PDFs slow and large. Cached until the file changes."""
    try:
        stamp = os.path.getmtime(logo_path)
    except OSError:
        stamp = None
    cached = _logo_cache.get(logo_path)
    if cached and cached[0] == stamp:
        return cached[1]

    with open(logo_path, "rb") as logo_file:
        raw = logo_file.read()
    try:
        from PIL import Image
        img = Image.open(io.BytesIO(raw))
        if img.height > LOGO_MAX_HEIGHT:
            img = img.resize((max(1, round(img.width * LOGO_MAX_HEIGHT / img.height)), LOGO_MAX_HEIGHT),
                             Image.LANCZOS)
            buf = io.BytesIO()
            img.save(buf, format="PNG", optimize=True)
            raw = buf.getvalue()
    except Exception:
        pass   # unreadable by Pillow - embed the original file
    encoded = base64.b64encode(raw).decode("utf-8")
    _logo_cache[logo_path] = (stamp, encoded)
    return encoded


# ==========================================================
# 🔹 PDF Report Generator
# ==========================================================

PDF_PAGE_CSS = """
    @page {
        size: A4;
        margin: 6mm;
    }

    body {
        transform: scale(0.86);
        transform-origin: top center;
    }
"""


def _report_details(df_string, df_cal_sum, batch_no):
    """Header values of one batch (Info + Summary rows)."""
    def get_value(k):
        hit = df_string.loc[df_string["Name"] == k, "Value"] if not df_string.empty else []
        return hit.iloc[0] if len(hit) else "N/A"

    def get_cal_value(k):
        hit = df_cal_sum.loc[df_cal_sum["Name"] == k, "Value"] if not df_cal_sum.empty else []
        return hit.iloc[0] if len(hit) else "N/A"

    return {
        "printed_date": datetime.now().strftime("%d-%m-%Y %H:%M"),
        "plant_name": get_value("Plant Name"),
        "recipe_name": get_value("Recipe Name"),
        "start_time": get_value("Start Date Time"),
        "end_time": get_value("End Date Time"),
        "shift": get_value("Shift"),
        "mixer_no": get_value("Mixer Selected"),
        "batch_no": batch_no,
        "time_taken": get_cal_value("BatchTimeMinutes"),
        "total_set_weight": get_cal_value("TotalBatchSetWeight"),
        "total_actual_weight": get_cal_value("TotalBatchActualWeight"),
    }


def _render_pdf(df_pivot, df_string, batch_no, df_cal_sum, include_speed=True, logo_base64=None):
    """One batch report, rendered (WeasyPrint document - pages can be merged)."""
    if logo_base64 is None:
        logo_base64 = encode_logo("data_files/logo.png")
    html = generate_html_report(
        df_pivot,
        logo_base64,
        _report_details(df_string, df_cal_sum, batch_no),
        include_speed=include_speed
    )
    return HTML(string=html).render(stylesheets=[CSS(string=PDF_PAGE_CSS)])


# ==========================================================
# 🔹 PDF Report Generator
# ==========================================================

def generate_pdf_report(
    df_pivot,
    df_string,
    batch_no,
    df_cal_sum,
    include_speed=True
):
    return _render_pdf(df_pivot, df_string, batch_no, df_cal_sum, include_speed).write_pdf()


def generate_multi_pdf_report(batches, include_speed=True):
    """batches: [(batch_no, df_pivot, df_string, df_cal_sum), ...] ->
    one PDF, every batch report starting on its own page."""
    logo_base64 = encode_logo("data_files/logo.png")
    docs = [_render_pdf(pivot, string, no, cal, include_speed, logo_base64)
            for no, pivot, string, cal in batches]
    pages = [page for doc in docs for page in doc.pages]
    return docs[0].copy(pages).write_pdf()


# ==========================================================
# 🔹 Excel Report Generator
# ==========================================================

def _fill_excel_sheet(ws, df_pivot, details, include_speed=True):
    """Writes one batch report onto worksheet ws."""
    # ======================================================
    # 🔹 Report Title
    # ======================================================
    ws.append(["BATCH REPORT"])
    ws.merge_cells("A1:F1")
    ws["A1"].font = Font(
        size=14,
        bold=True
    )
    ws["A1"].alignment = Alignment(
     horizontal="center"
    )

    # ======================================================
    # 🔹 Report Information
    # ======================================================
    ws.append([
        "Printed Date:",
        details["printed_date"]
    ])
    ws.append([])
    headers = [
        ("Plant Name",details["plant_name"]),
        ("Recipe Name",details["recipe_name"]),
        ("Batch No",details["batch_no"]),
        ("Mixer No",details["mixer_no"]),
        ("Start Time",details["start_time"]),
        ("End Time",details["end_time"]),
        ("Shift",details["shift"]),
        ("Total Set Weight (Kg)",
        df_pivot["SetWeight"].sum()),
        ("Total Actual Weight (Kg)",round(df_pivot["ActualWeight"].sum(),2))
    ]
    for i in range(
        0,
        len(headers),
        2
    ):
        row = [
            headers[i][0],
            headers[i][1]
        ]
        if i + 1 < len(headers):
            row += [
                headers[i + 1][0],
                headers[i + 1][1]
            ]
        ws.append(row)
    ws.append([])
    # ======================================================
    # 🔹 Table Headers
    # ======================================================

    table_headers = [
        "Silo No",
        "Material Name",
        "Set Weight",
        "Actual Weight",
        "Difference",
        "Tolerance"
    ]

    # Add speed columns only when enabled
    if include_speed:
        table_headers += [
            "CoarseSpeed",
            "FineSpeed"
        ]
    table_headers += [
        "Start Time",
        "End Time",
        "Duration (hh:mm:ss)"
    ]
    ws.append(table_headers)

    # ======================================================
    # 🔹 Header Formatting
    # ======================================================

    header_fill = PatternFill(
        start_color="DDDDDD",
        fill_type="solid"
    )
    for col in range(1,len(table_headers) + 1):
        cell = ws.cell(
            row=ws.max_row,
            column=col
        )
        cell.font = Font(
            bold=True
        )
        cell.alignment = Alignment(
            horizontal="center"
        )
        cell.fill = header_fill

    # ======================================================
    # 🔹 Add Data Rows
    # ======================================================
    for _, row in df_pivot.iterrows():

        row_data = [
        row.get("SiloNo",""),
        row.get("MaterialName",""),
        row.get("SetWeight",""),
        row.get("ActualWeight",""),
        row.get("Difference",""),
        row.get("Tolerance","")
        ]

        # Add speed values only when enabled
        if include_speed:
            row_data += [
                row.get(
                    "CoarseSpeed",
                    ""
                ),
                row.get(
                    "FineSpeed",
                    ""
                )
            ]
        row_data += [
            row.get("StartTime", ""),
            row.get("EndTime", ""),
            row.get("Duration", "")
        ]
        ws.append(row_data)



def generate_excel_report(
    df_pivot,
    df_string,
    batch_no,
    df_cal_sum,
    include_speed=True
):
    output = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Batch Report"
    _fill_excel_sheet(ws, df_pivot, _report_details(df_string, df_cal_sum, batch_no), include_speed)
    wb.save(output)
    output.seek(0)
    return output.read()


def generate_multi_excel_report(batches, summary=None, include_speed=True):
    """One workbook: a "Summary" sheet (one row per batch) and one sheet per
    batch with the same layout as the single batch report."""
    output = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    if summary is not None and not summary.empty:
        ws.append(list(summary.columns))
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="DDDDDD", fill_type="solid")
        for row in summary.itertuples(index=False):
            ws.append(list(row))
        for i, col in enumerate(summary.columns, start=1):
            width = max([len(str(col))] + [len(str(v)) for v in summary.iloc[:, i - 1].head(300)]) + 2
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(width, 40)
        ws.freeze_panes = "A2"

    for no, pivot, string, cal in batches:
        sheet = wb.create_sheet(title=f"Batch {no}"[:31])
        _fill_excel_sheet(sheet, pivot, _report_details(string, cal, no), include_speed)

    wb.save(output)
    output.seek(0)
    return output.read()

# ==========================================================
# 🔹 HTML Template for PDF
# ==========================================================

def generate_html_report(
    df,
    logo_base64,
    details,
    include_speed=True
):

    # ======================================================
    # 🔹 Generate Data Rows
    # ======================================================

    data_rows = ""
    for _, row in df.iterrows():
        data_rows += f"""
        <tr>
            <td>{row["SiloNo"]}</td>
            <td>{row["MaterialName"]}</td>
            <td>{row["SetWeight"]}</td>
            <td>{row["ActualWeight"]}</td>
            <td>{row["Difference"]}</td>
            <td>{row["Tolerance"]}</td>
        """

        # Add speed columns only when enabled
        if include_speed:
            data_rows += f"""
            <td>{row.get("CoarseSpeed", "")}</td>
            <td>{row.get("FineSpeed", "")}</td>
            """
        data_rows += f"""
            <td>{row.get("StartTime", "")}</td>
            <td>{row.get("EndTime", "")}</td>
            <td>{row.get("Duration", "")}</td>
        </tr>
        """

    # ======================================================
    # 🔹 Speed Table Headers
    # ======================================================
    speed_headers = ""
    if include_speed:
        speed_headers = """
            <th>Coarse Speed (%)</th>
            <th>Fine Speed (%)</th>
        """

    # ======================================================
    # 🔹 HTML Template
    # ======================================================
    html_template = f"""
    <!DOCTYPE html>
    <html>
    <head>
        <title>Batch Report</title>
        <style>
            @import url(
                'https://fonts.googleapis.com/css2?family=Roboto:wght@400;700&display=swap'
            );
            @page {{
                size: A4;
                margin: 20mm;
            }}
            body {{
                font-family: 'Roboto', sans-serif;
                margin: 20px;
                font-size: 12pt;
            }}
            h1 {{
                text-align: center;
                margin-bottom: 20px;
            }}
            .header {{
                display: flex;
                justify-content: space-between;
                align-items: center;
                margin-bottom: 20px;
            }}
            .printed-date {{
                text-align: right;
                font-weight: bold;
                font-size: 12pt;
            }}
            .container {{
                margin-bottom: 20px;
            }}
            .info-section table {{
                width: 100%;
                border-collapse: collapse;
                margin-bottom: 20px;
            }}
            .info-section td {{
                padding: 8px;
                text-align: left;
                border: none;
            }}
            table {{
                width: 100%;
                border-collapse: collapse;
                margin-bottom: 20px;
            }}
            td,
            th {{
                padding: 8px 6px;
                text-align: center;
                border: 1px solid #ddd;
                font-size: 10pt;
            }}
            th {{
                background-color: #f2f2f2;
                font-weight: bold;
            }}
            .footer {{
                text-align: center;
                margin-top: 20px;
                font-weight: bold;
            }}
            .small-footer {{
                text-align: center;
                font-size: 10px;
                margin-top: 10px;
            }}
        </style>
    </head>
    <body>
        <h1>BATCH REPORT</h1>
        <div class="header">
            <img
                src="data:image/png;base64,{logo_base64}"
                alt="Logo"
                style="width: 100px;"
            >
            <div class="printed-date">
                Printed Date: {details['printed_date']}
            </div>
        </div>
        <div class="container">
            <div class="info-section">
                <table>
                    <tr>
                        <td><b>Recipe Name:</b></td>
                        <td>{details['recipe_name']}</td>

                        <td><b>Time Taken(Min):</b></td>
                        <td>{details['time_taken']}</td>
                    </tr>

                    <tr>
                        <td><b>Batch No:</b></td>
                        <td>{details['batch_no']}</td>

                        <td><b>Total Set Weight(Kg):</b></td>
                        <td>{details['total_set_weight']} Kg</td>
                    </tr>
                    <tr>
                        <td><b>Start Time:</b></td>
                        <td>{details['start_time']}</td>
                        
                        <td><b>Total Actual Weight(Kg):</b></td>
                        <td>{details['total_actual_weight']} Kg</td>
                    </tr>

                    <tr>
                        <td><b>End Time:</b></td>
                        <td>{details['end_time']}</td>

                        <td><b>Shift:</b></td>
                        <td>{details['shift']}</td>
                    </tr>

                    <tr>
                        <td><b>Mixer No:</b></td>
                        <td>{details['mixer_no']}</td>

                        <td><b>Plant Name:</b></td>
                        <td>{details['plant_name']}</td>
                    </tr>
                </table>
            </div>
        </div>
        <h2>Data Table</h2>
        <table>
            <tr>
                <th>Silo No</th>
                <th>Material Name</th>
                <th>Set Weight (Kg)</th>
                <th>Actual Weight (Kg)</th>
                <th>Difference (Kg)</th>
                <th>Tolerance (Kg)</th>
                {speed_headers}
                <th>Start Time</th>
                <th>End Time</th>
                <th>Duration (hh:mm:ss)</th>
            </tr>
            {data_rows}
        </table>
        <div class="small-footer">
            This report is generated by Skew Reporting Software,
            developed by Prolite Automation.
        </div>
    </body>
    </html>
    """
    return html_template